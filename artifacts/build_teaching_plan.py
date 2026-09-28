from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = r'C:\chess\artifacts\Flutter Firebase Chess Teaching Plan.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.7)
sec.right_margin = Inches(.7)

for style_name, size in [('Normal', 10.5), ('Title', 24), ('Heading 1', 17), ('Heading 2', 13)]:
    s = doc.styles[style_name]
    s.font.name = 'Aptos'
    s._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].paragraph_format.space_after = Pt(5)

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    props.append(shd)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.width = Inches(widths[i])
        shade(c, '243B53')
        for r in c.paragraphs[0].runs:
            r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(9)
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].width = Inches(widths[i])
            if index % 2: shade(cells[i], 'F1F5F9')
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size = Pt(8.8)
    doc.add_paragraph()

def bullets(items):
    for item in items: doc.add_paragraph(item, style='List Bullet')

def code(text):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'
    c = t.cell(0, 0); c.text = text; shade(c, 'F3F4F6')
    for r in c.paragraphs[0].runs:
        r.font.name = 'Consolas'; r._element.rPr.rFonts.set(qn('w:ascii'), 'Consolas'); r.font.size = Pt(8.4)
    doc.add_paragraph()

def section(day, title, outcome, flow, theory, code_text, walkthrough, activity, homework):
    doc.add_heading(f'Day {day}  {title}', 1)
    p = doc.add_paragraph(); p.add_run('End of session outcome: ').bold = True; p.add_run(outcome)
    doc.add_heading('One Hour Flow', 2)
    table(['Time', 'Block', 'Teacher actions', 'Student actions'], flow, [.6, 1.25, 2.85, 1.85])
    doc.add_heading('Theory Concepts to Teach', 2); bullets(theory)
    doc.add_heading('Code Explanation Block', 2); code(code_text)
    table(['File or concept', 'What to explain', 'Question to ask'], walkthrough, [1.75, 3.2, 1.6])
    doc.add_heading('Guided Activity', 2); doc.add_paragraph(activity)
    doc.add_heading('Homework or Exit Ticket', 2); doc.add_paragraph(homework)

p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Six Day Flutter Firebase Chess Teaching Plan')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('One hour sessions for students who know Flutter basics').italic = True
doc.add_heading('Purpose', 1)
doc.add_paragraph('This guide uses a multiplayer chess app to teach Firebase backend ideas. Students already know Flutter screens and widgets from previous projects. The teaching emphasis is identity, cloud data, real-time updates, safety, testing, and deployment.')
doc.add_heading('Preparation Before Day 1', 2)
bullets(['Enable Anonymous Authentication in Firebase Authentication.', 'Create Cloud Firestore and confirm the app can write a games document.', 'Deploy the app and open it in two browsers or on two devices.', 'Keep Firebase Console open: Authentication, Firestore Data, Rules, and Hosting.'])
doc.add_heading('Core Vocabulary', 2)
table(['Term', 'Student friendly meaning', 'Chess example'], [
    ('Authentication', 'Giving a user an identity.', 'Anonymous sign-in creates a player uid.'),
    ('Firestore document', 'A cloud record with named fields.', 'One match is one games document.'),
    ('Stream', 'Data that can update repeatedly.', 'Opponent move updates the board live.'),
    ('FEN', 'Text that describes a chess position.', 'The board reloads from this field.'),
    ('Transaction', 'Read, check, then safely write.', 'Only one user can claim Black.'),
    ('Security rule', 'Firebase permission enforced on the backend.', 'Only players should update a game.'),
], [1.35, 2.65, 2.55])

section(1, 'From Screen to Shared Game',
'Students can trace one chess move from a tap on one device to a board update on another device.',
[('0-10', 'Demo', 'Create and join a game in two browsers.', 'Observe both boards.'), ('10-22', 'Architecture', 'Draw UI to Provider to Service to Firestore to Stream.', 'Label project files.'), ('22-40', 'Code tour', 'Open main.dart, providers, and game service.', 'Find create, join, listen, and move methods.'), ('40-52', 'Console', 'Inspect a games document in Firestore.', 'Match fields to UI.'), ('52-60', 'Review', 'Ask prediction questions.', 'Explain the data flow.')],
['Frontend is what the player sees; backend stores shared data and enforces permissions.', 'Two devices share data, not widgets or pixels.', 'Provider is local Flutter state; Firestore is shared cloud state.', 'A real-time stream reacts to changes without a manual refresh.'],
'Flutter UI -> ChessProvider -> GameService -> Firestore\nFirestore snapshot stream -> ChessProvider -> Flutter UI on both devices',
[('main.dart', 'Firebase initializes before the widget tree. MultiProvider exposes shared state.', 'Why must Firebase initialize first?'), ('game_service.dart', 'Service methods are the boundary between Flutter and Firestore.', 'Which method creates a document?'), ('chess_provider.dart', 'The provider turns document data into board state and calls notifyListeners.', 'Why does the UI rebuild?')],
'In pairs, make one move as White and use the Firestore console to locate the changed fen, pgn, and turn fields.',
'Draw the architecture with arrows. Write one responsibility for main.dart, AuthProvider, ChessProvider, GameService, LobbyScreen, and GameScreen.')

section(2, 'Giving Every Player an Identity',
'Students can enable Anonymous Authentication and explain why a game needs a unique player ID.',
[('0-8', 'Recap', 'Ask how Firebase knows who clicked Create Game.', 'Answer from the architecture map.'), ('8-20', 'Console setup', 'Show Authentication and enable Anonymous provider.', 'Follow on their project or observe.'), ('20-36', 'Code reading', 'Explain AuthProvider and LoginScreen.', 'Locate sign-in and state listener.'), ('36-50', 'Build', 'Add a player ID or nickname label.', 'Implement and test.'), ('50-60', 'Failures', 'Show a disabled provider error.', 'Translate the error into a fix.')],
['Authentication asks who is this user; authorization asks what can this user do.', 'Anonymous authentication still creates a genuine Firebase user with a uid.', 'authStateChanges is a stream that emits sign-in and sign-out events.', 'Network calls are Futures and can fail, so await and error handling matter.'],
"_auth.authStateChanges().listen((User? user) {\n  _user = user;\n  notifyListeners();\n});\n\nawait _auth.signInAnonymously();",
[('AuthProvider', 'The listener saves the current User and asks consumers to rebuild.', 'What does null user mean?'), ('LoginScreen', 'The button waits for Firebase and now displays sign-in failures.', 'What happens if Anonymous Auth is disabled?'), ('uid', 'The uid is saved as whitePlayerId or blackPlayerId.', 'Why should a display name not be used as identity?')],
'Add a sign-out action in the lobby. Verify the auth stream returns the app to LoginScreen without manually navigating there.',
'Compare Anonymous Authentication with email-password authentication. Give one advantage and one limitation of each.')

section(3, 'Saving a Chess Match in Firestore',
'Students can model a game as a Firestore document and relate every field to a visible game behaviour.',
[('0-10', 'Database basics', 'Explain collection, document, field, and JSON-like data.', 'Sketch a games collection.'), ('10-25', 'Data model', 'Read GameModel toFirestore and fromFirestore.', 'List required match fields.'), ('25-42', 'Create and join', 'Trace createGame and joinGame.', 'Create and join in two browsers.'), ('42-54', 'Build', 'Add createdAt or game title.', 'Store and show the field.'), ('54-60', 'Review', 'Contrast local variables with Firestore.', 'Say what survives refresh.')],
['Firestore is a NoSQL database organised as collections and documents.', 'A document is the shared source of truth for one game.', 'FEN restores a precise board position; PGN is readable move history.', 'A document ID is the shareable code a second player uses to join.'],
"final game = GameModel(\n  id: '', fen: chess.Chess().fen, pgn: '', turn: 'w',\n  whitePlayerId: userId, status: 'waiting',\n);\nfinal ref = await _db.collection('games').add(game.toFirestore());",
[('fen', 'Contains pieces, side to move, castling and counters.', 'Why is a board screenshot not enough?'), ('pgn', 'Contains readable chess notation.', 'Why keep history separately from FEN?'), ('status', 'waiting, active, or finished drives the UI.', 'What should the lobby show for waiting?'), ('player IDs', 'They connect Firebase users to chess colours.', 'How do we know who may move?')],
'Students create a game, inspect its document, then add a gameName field and render it in the lobby.',
'Propose three useful fields for a future chess feature and describe the UI each one enables.')

section(4, 'Real Time Synchronization and State',
'Students can explain snapshots, streams, provider rebuilds, and why both devices agree on the board.',
[('0-8', 'Experiment', 'Make a move in one window.', 'Narrate what the second window does.'), ('8-22', 'Theory', 'Compare function return, Future, and Stream.', 'Predict how many times each produces data.'), ('22-40', 'Trace code', 'Read streamGame and initGame.', 'Follow a snapshot into the chess engine.'), ('40-52', 'Build', 'Improve waiting, turn, or check status.', 'Test all status values.'), ('52-60', 'Debug', 'Use an invalid game ID.', 'Find the source and suggest a message.')],
['A Future gives one later value; a Stream gives values repeatedly.', 'A Firestore snapshot is the current version of a document.', 'The listener receives the first snapshot and every later update.', 'Local state keeps the screen responsive; shared state keeps devices consistent.'],
"Stream<GameModel> streamGame(String gameId) {\n  return _db.collection('games').doc(gameId).snapshots()\n      .map((doc) => GameModel.fromFirestore(doc));\n}\n\n_subscription = _gameService.streamGame(gameId).listen((gameModel) {\n  _status = gameModel.status;\n  notifyListeners();\n});",
[('snapshots()', 'Starts a live document listener.', 'Why is this not the same as get()?'), ('map()', 'Converts raw Firestore data into GameModel.', 'Why use a model instead of maps everywhere?'), ('notifyListeners()', 'Rebuilds Consumer widgets that use the provider.', 'What if it is omitted?')],
'Change one field in Firestore during a controlled demo and ask students to predict exactly which widget changes. Then restore it.',
'Write a six-step sequence for a Black move using Provider, Firestore, document, stream, and rebuild.')

section(5, 'Protecting Games with Rules and Transactions',
'Students can distinguish UI checks from backend enforcement and explain how transactions prevent duplicate or stale actions.',
[('0-10', 'Threat model', 'Ask whether hiding a button is security.', 'Explain why modified clients exist.'), ('10-25', 'Rules', 'Teach request.auth and document checks.', 'Read a simple classroom rule.'), ('25-42', 'Transactions', 'Trace joinGame and submitMove.', 'Identify each validation condition.'), ('42-54', 'Build', 'Improve full-game and self-join errors.', 'Try invalid joins.'), ('54-60', 'Exit ticket', 'Ask one rules and one transaction question.', 'Answer in one sentence.')],
['Client validation is for user experience; Security Rules are for backend protection.', 'Transactions read the latest state, validate it, and write only if the state is still valid.', 'The join transaction prevents two users from both taking Black.', 'The move transaction checks identity, turn, status, and previous FEN to avoid stale writes.'],
"await _db.runTransaction((transaction) async {\n  final snapshot = await transaction.get(ref);\n  final game = GameModel.fromFirestore(snapshot);\n  if (game.status != 'waiting' || game.blackPlayerId != null) {\n    throw StateError('This game already has two players.');\n  }\n  transaction.update(ref, {'blackPlayerId': userId, 'status': 'active'});\n});",
[('request.auth', 'The Firebase user making a request.', 'Why is it safer than trusting a uid sent by the app?'), ('resource.data', 'The existing document in Firestore.', 'What does it tell a join transaction?'), ('request.resource.data', 'The proposed new document.', 'What could a rule compare before accepting it?')],
'Have students try to join the same waiting game at nearly the same time. Discuss why only one should win the Black seat.',
'Write pseudocode for a rule that lets a player update their own game but prevents a random signed-in user from changing fen.')

section(6, 'Test Deploy and Extend',
'Students can verify multiplayer behaviour, deploy the application, and plan a backend-aware feature extension.',
[('0-10', 'Review', 'Rebuild the full data flow from memory.', 'Name the role of every backend component.'), ('10-22', 'Testing', 'Introduce a two-player checklist.', 'Run it in pairs.'), ('22-42', 'Feature sprint', 'Offer one small extension.', 'Implement a scoped feature.'), ('42-52', 'Deploy', 'Build web and deploy Firebase Hosting.', 'Open the hosted build incognito.'), ('52-60', 'Present', 'Invite short demos and reflection.', 'Show one feature and one backend lesson.')],
['Testing describes expected behaviour, especially between two users.', 'Deployment is separate from local running: build output must be uploaded to Hosting.', 'New features begin with data design: fields, permissions, stream effect, and UI.', 'A feature is not complete until it works after refresh and handles invalid actions.'],
"flutter analyze\nflutter test\nflutter build web\nfirebase deploy --only hosting",
[('Two browser test', 'Test create, join, both colours, resign, rematch, and refresh.', 'What should happen if the second player refreshes?'), ('Feature design', 'Choose fields before building UI.', 'Who owns this field and who may change it?'), ('Deployment', 'Build sends Flutter code to build/web; Firebase deploy publishes it.', 'Why can an old website remain after local code changes?')],
'Pick one: nicknames, draw offer, rematch approval, chess clock, or lobby list. First write the document fields and permissions, then code a small UI slice.',
'Complete the test checklist and submit a feature proposal with data fields, allowed writers, UI change, and one edge case.')

doc.add_heading('Two Player Test Checklist', 1)
table(['Scenario', 'Expected result', 'Check'], [
    ('Sign in', 'Each browser gets a different anonymous uid.', '____'), ('Create', 'White receives a game ID and waits.', '____'), ('Join', 'One second user becomes Black; game becomes active.', '____'), ('Move', 'Both boards update and turn switches.', '____'), ('Illegal move', 'No document position change occurs.', '____'), ('Resign', 'Both players see a result and rematch action.', '____'), ('Refresh', 'Current position reloads from Firestore.', '____'), ('Deploy', 'Hosted app shows the latest build after hard refresh.', '____'),
], [1.6, 4.7, .6])
doc.add_heading('Feature Sprint Menu', 1)
table(['Feature', 'Suggested Firestore fields', 'Theory question'], [
    ('Nicknames', 'whiteName, blackName', 'Who can edit a name after the game begins?'), ('Draw offer', 'drawOfferedBy, drawAccepted', 'How do two users agree on one result?'), ('Rematch approval', 'whiteRematchReady, blackRematchReady', 'When is it safe to reset the board?'), ('Chess clock', 'whiteTimeMs, blackTimeMs, lastMoveAt', 'Which clock data can clients trust?'), ('Lobby list', 'Query waiting games', 'Which users should be allowed to read it?'),
], [1.45, 2.75, 2.7])
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = footer.add_run('Flutter Firebase Chess Teaching Plan')
rr.font.size = Pt(8); rr.font.color.rgb = RGBColor(100, 100, 100)
doc.save(OUTPUT)
print(OUTPUT)
